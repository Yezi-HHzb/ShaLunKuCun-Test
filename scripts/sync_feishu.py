import argparse
import hashlib
import json
import os
import re
import tempfile
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU


DEFAULT_TEMPLATE = "E10品号-规格-30B套圈.xlsx"
DEFAULT_OUTPUT = "E10库存二维码.xlsx"
DEFAULT_DATA = "data.json"
DEFAULT_MANIFEST = "qr_manifest.json"
DEFAULT_QR_DIR = "qrcodes"
DEFAULT_SITE_URL = "https://yezi-hhzb.github.io/ShaLunKuCun-Test/"
SHEET_NAME = "最新"
PART_COLUMN = "F"
QR_COLUMN = "J"
MISSING_TEXT = "E10库存二维码"
HTTP_TIMEOUT = 30


class FeishuError(RuntimeError):
    pass


def utc_now():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def extract_value(raw):
    if isinstance(raw, list):
        return " ".join(filter(None, (extract_value(item) for item in raw)))
    if isinstance(raw, dict):
        for key in ("text", "number", "value"):
            if key in raw:
                return str(raw[key]).strip()
        return ""
    return str(raw).strip() if raw is not None else ""


def normalize_part_no(value):
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_records(records):
    result = []
    for rec in records:
        fields = rec.get("fields", {})
        part_no = normalize_part_no(
            extract_value(fields.get("厂料号", ""))
            or extract_value(fields.get("品号", ""))
        )
        if not part_no:
            continue
        result.append(
            {
                "partNo": part_no,
                "spec": extract_value(fields.get("规格型号", "")),
                "manufacturer": extract_value(fields.get("厂家", "")),
                "stock": extract_value(fields.get("库存", "")),
                "linkedStock": extract_value(fields.get("联库", "")),
                "totalStock": extract_value(fields.get("总库", "")),
                "backupStock": extract_value(fields.get("备货", "")),
                "process": extract_value(fields.get("工序", "")),
            }
        )
    return result


def _post_json(session, url, request_exception, **kwargs):
    try:
        response = session.post(url, timeout=HTTP_TIMEOUT, **kwargs)
        response.raise_for_status()
        payload = response.json()
    except (request_exception, ValueError) as exc:
        raise FeishuError(f"飞书请求失败: {exc}") from exc
    if payload.get("code") != 0:
        raise FeishuError(f"飞书接口返回错误: {payload}")
    return payload


def fetch_feishu_records():
    import requests

    required = (
        "FEISHU_APP_ID",
        "FEISHU_APP_SECRET",
        "FEISHU_APP_TOKEN",
        "FEISHU_TABLE_ID",
    )
    missing = [name for name in required if not os.environ.get(name)]
    if missing:
        raise FeishuError("缺少环境变量: " + ", ".join(missing))

    session = requests.Session()
    token_payload = _post_json(
        session,
        "https://open.feishu.cn/open-apis/auth/v3/tenant_access_token/internal",
        requests.RequestException,
        json={
            "app_id": os.environ["FEISHU_APP_ID"],
            "app_secret": os.environ["FEISHU_APP_SECRET"],
        },
    )
    token = token_payload["tenant_access_token"]
    url = (
        "https://open.feishu.cn/open-apis/bitable/v1/apps/"
        f"{os.environ['FEISHU_APP_TOKEN']}/tables/"
        f"{os.environ['FEISHU_TABLE_ID']}/records/search"
    )
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}

    records = []
    page_token = None
    while True:
        body = {"page_size": 500}
        if page_token:
            body["page_token"] = page_token
        payload = _post_json(session, url, requests.RequestException, headers=headers, json=body)
        data = payload.get("data", {})
        items = data.get("items", [])
        if not isinstance(items, list):
            raise FeishuError("飞书记录格式异常，已停止更新文件")
        records.extend(items)
        if not data.get("has_more"):
            break
        page_token = data.get("page_token")
        if not page_token:
            raise FeishuError("飞书分页信息异常，已停止更新文件")
    return records


def load_json(path, fallback):
    path = Path(path)
    if not path.exists():
        return fallback
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json_atomic(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", encoding="utf-8", dir=path.parent, delete=False, suffix=".tmp"
    ) as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        temp_name = handle.name
    os.replace(temp_name, path)


def qr_filename(part_no):
    readable = re.sub(r"[^0-9A-Za-z._-]+", "_", part_no).strip("._-")[:48]
    readable = readable or "part"
    digest = hashlib.sha256(part_no.encode("utf-8")).hexdigest()[:12]
    return f"{readable}-{digest}.png"


def qr_url(site_url, part_no):
    return f"{site_url.rstrip('/')}/?partNo={quote(part_no, safe='')}"


def generate_qr(path, content):
    import qrcode

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    code = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=4)
    code.add_data(content)
    code.make(fit=True)
    code.make_image(fill_color="black", back_color="white").save(path)


def update_manifest(data, manifest_path, qr_dir, site_url, now):
    manifest = load_json(manifest_path, {"version": 1, "items": {}})
    items = manifest.setdefault("items", {})
    active_parts = {normalize_part_no(item.get("partNo")) for item in data}
    active_parts.discard("")

    for part_no in sorted(active_parts):
        content = qr_url(site_url, part_no)
        entry = items.get(part_no, {})
        filename = entry.get("file") or str(Path(qr_dir) / qr_filename(part_no)).replace("\\", "/")
        if entry.get("url") != content or not Path(filename).exists():
            generate_qr(filename, content)
        entry.update(
            {
                "partNo": part_no,
                "file": filename,
                "url": content,
                "status": "active",
                "firstSeenAt": entry.get("firstSeenAt", now),
                "lastSeenAt": now,
                "retiredAt": None,
            }
        )
        items[part_no] = entry

    for part_no, entry in items.items():
        if part_no not in active_parts and entry.get("status") != "inactive":
            entry["status"] = "inactive"
            entry["retiredAt"] = now

    manifest["updatedAt"] = now
    return manifest


def _column_pixels(sheet, column):
    width = sheet.column_dimensions[column].width or 13
    return int(width * 7 + 5)


def _row_pixels(sheet, row):
    points = sheet.row_dimensions[row].height or 15
    return int(points * 96 / 72)


def _place_image(sheet, image_path, row, size=170):
    image = ExcelImage(image_path)
    image.width = size
    image.height = size
    col_index = sheet[QR_COLUMN + "1"].column - 1
    col_offset = max(0, (_column_pixels(sheet, QR_COLUMN) - size) // 2)
    row_offset = max(0, (_row_pixels(sheet, row) - size) // 2)
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=col_index,
            colOff=pixels_to_EMU(col_offset),
            row=row - 1,
            rowOff=pixels_to_EMU(row_offset),
        ),
        ext=XDRPositiveSize2D(cx=pixels_to_EMU(size), cy=pixels_to_EMU(size)),
    )
    sheet.add_image(image)


def generate_excel(template_path, output_path, manifest):
    workbook = load_workbook(template_path)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"模板缺少工作表: {SHEET_NAME}")
    sheet = workbook[SHEET_NAME]

    qr_col_index = sheet[QR_COLUMN + "1"].column - 1
    sheet._images = [
        image
        for image in sheet._images
        if not (
            hasattr(image.anchor, "_from")
            and image.anchor._from.col == qr_col_index
            and image.anchor._from.row >= 2
        )
    ]

    active = {
        part_no: entry
        for part_no, entry in manifest.get("items", {}).items()
        if entry.get("status") == "active"
    }
    for row in range(3, sheet.max_row + 1):
        part_no = normalize_part_no(sheet[f"{PART_COLUMN}{row}"].value)
        if not part_no or part_no == "品号":
            continue
        target = sheet[f"{QR_COLUMN}{row}"]
        entry = active.get(part_no)
        if entry and Path(entry["file"]).exists():
            target.value = None
            _place_image(sheet, entry["file"], row)
        else:
            target.value = MISSING_TEXT
            alignment = copy(target.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            alignment.wrap_text = True
            target.alignment = alignment

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args():
    parser = argparse.ArgumentParser(description="同步飞书库存并生成二维码 Excel")
    parser.add_argument("--input-json", help="本地测试用的飞书 records JSON；不调用飞书 API")
    parser.add_argument("--template", default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--data", default=DEFAULT_DATA)
    parser.add_argument("--manifest", default=DEFAULT_MANIFEST)
    parser.add_argument("--qr-dir", default=DEFAULT_QR_DIR)
    parser.add_argument("--site-url", default=os.environ.get("SITE_BASE_URL", DEFAULT_SITE_URL))
    return parser.parse_args()


def main():
    args = parse_args()
    if args.input_json:
        records = load_json(args.input_json, [])
    else:
        records = fetch_feishu_records()
    if not records and os.environ.get("ALLOW_EMPTY_DATA") != "1":
        raise FeishuError("飞书返回 0 条记录。为防止误清空，已停止更新文件")
    data = parse_records(records)
    now = utc_now()
    manifest = update_manifest(data, args.manifest, args.qr_dir, args.site_url, now)
    generate_excel(args.template, args.output, manifest)
    write_json_atomic(args.data, data)
    write_json_atomic(args.manifest, manifest)
    print(f"成功同步 {len(data)} 条记录，生成 {args.output}")


if __name__ == "__main__":
    main()
